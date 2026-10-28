"""Fill own/<VENUE>/<TICKER>.png for every instrument the nvstly icon set does not cover.

Commands
  build        offline: copy nvstly logos of the company's US listing / render simple-icons glyphs
  tradingview  online:  download each company's TradingView logo (run this on your own PC);
               replaces everything not already from TradingView
  placeholders last resort: lettered circle for any ticker that still has no file
  report       write reports/logo_status.csv + reports/logo_status.md

Every PNG written follows README.txt: 250x250, 8-bit RGBA, not interlaced, < 256 KB.
own/sources.csv records where each file came from.
"""
import argparse
import csv
import io
import json
import os
import re
import sys
import urllib.parse
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCES = os.path.join(ROOT, "tools", "logo_sources.csv")
NVSTLY_INDEX = os.path.join(ROOT, "tools", "nvstly_tickers.txt")
OWN = os.path.join(ROOT, "own")
US_VENUES = {"NYQ", "NMS", "ASE"}
SIZE = 250


def load_sources():
    with open(SOURCES, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def own_path(own, venue, symbol):
    return os.path.join(own, venue, symbol.replace(":", "_") + ".png")


def load_log(own):
    path = os.path.join(own, "sources.csv")
    if not os.path.exists(path):
        return {}
    with open(path, newline="", encoding="utf-8") as f:
        return {(r["venue"], r["symbol"]): r["source"] for r in csv.DictReader(f)}


def save_log(own, log):
    with open(os.path.join(own, "sources.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["venue", "symbol", "source"])
        for (venue, symbol), source in sorted(log.items()):
            w.writerow([venue, symbol, source])


def save_png(img, path):
    """Fit img on a transparent 250x250 canvas and save it within the README rules."""
    from PIL import Image

    img = img.convert("RGBA")
    bbox = img.getchannel("A").getbbox()
    if bbox:
        img = img.crop(bbox)
    scale = SIZE / max(img.size)
    img = img.resize((max(1, round(img.width * scale)), max(1, round(img.height * scale))), Image.LANCZOS)
    canvas = Image.new("RGBA", (SIZE, SIZE), (0, 0, 0, 0))
    canvas.paste(img, ((SIZE - img.width) // 2, (SIZE - img.height) // 2), img)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    canvas.save(path, "PNG", optimize=True)
    if os.path.getsize(path) > 256 * 1024:
        raise ValueError(f"{path} is larger than 256 KB")


# ---------------------------------------------------------------- build (offline)

def render_simple_icon(si_dir, title, size=SIZE):
    import cairosvg
    from PIL import Image

    data = json.load(open(os.path.join(si_dir, "data", "simple-icons.json"), encoding="utf-8"))
    entry = next(d for d in data if d["title"] == title)
    slug = entry.get("slug") or slugify(title)
    svg = open(os.path.join(si_dir, "icons", slug + ".svg"), encoding="utf-8").read()
    color = "#" + entry["hex"]
    r, g, b = (int(entry["hex"][i:i + 2], 16) for i in (0, 2, 4))
    if 0.2126 * r + 0.7152 * g + 0.0722 * b < 60:  # too dark for a dark UI
        color = "#FFFFFF"
    svg = svg.replace("<svg ", f'<svg fill="{color}" ', 1)
    png = cairosvg.svg2png(bytestring=svg.encode(), output_width=size * 2, output_height=size * 2)
    return Image.open(io.BytesIO(png))


def slugify(title):
    """simple-icons' titleToSlug."""
    import unicodedata

    s = title.lower().replace("+", "plus").replace(".", "dot").replace("&", "and")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", s)


def cmd_build(args):
    from PIL import Image

    log = load_log(args.own)
    nv_dir = os.path.join(args.nvstly, "ticker_icons")
    for row in load_sources():
        key = (row["venue"], row["symbol"])
        path = own_path(args.own, *key)
        if os.path.exists(path) and not args.overwrite:
            continue
        alias = next((a for a in row["nvstly_aliases"].split(";")
                      if a and os.path.exists(os.path.join(nv_dir, a + ".png"))), None)
        if alias:
            save_png(Image.open(os.path.join(nv_dir, alias + ".png")), path)
            log[key] = f"nvstly:{alias}"
        elif row["simple_icons"]:
            save_png(render_simple_icon(args.simple_icons, row["simple_icons"]), path)
            log[key] = f"simple-icons:{row['simple_icons']}"
    save_log(args.own, log)
    print(f"{len(log)} logos in {args.own}")


# ---------------------------------------------------------------- tradingview (online)

SCAN_URL = "https://scanner.tradingview.com/global/scan"
SEARCH_URL = "https://symbol-search.tradingview.com/symbol_search/v3/?"
LOGO_URL = "https://s3-symbol-logo.tradingview.com/{}--big.svg"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/130.0 Safari/537.36"}


def tradingview_logoids(tickers):
    """Logo ids from TradingView's screener, e.g. {'TSE:6857': 'advantest'}."""
    body = json.dumps({"symbols": {"tickers": tickers}, "columns": ["logoid", "description"]}).encode()
    req = urllib.request.Request(SCAN_URL, data=body, headers={**UA, "Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            data = json.load(r)
    except Exception as e:
        print(f"  screener lookup failed ({e}); falling back to symbol search")
        return {}
    return {d["s"]: d["d"][0] for d in data.get("data", []) if d["d"] and d["d"][0]}


def tradingview_search_logoid(ticker):
    """Fallback: TradingView's symbol search, as used by its own search box."""
    exchange, symbol = ticker.split(":", 1)
    url = SEARCH_URL + urllib.parse.urlencode({"text": symbol, "exchange": exchange, "lang": "en",
                                               "domain": "production", "search_type": "undefined"})
    req = urllib.request.Request(url, headers={**UA, "Origin": "https://www.tradingview.com",
                                                 "Referer": "https://www.tradingview.com/"})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            data = json.load(r)
    except Exception:
        return None
    for hit in data.get("symbols", []) if isinstance(data, dict) else data:
        name = re.sub(r"</?em>", "", hit.get("symbol", ""))
        if name.upper() == symbol.upper() and hit.get("logoid"):
            return hit["logoid"]
    return None


def render_svg(page, svg, square=False):
    """Rasterise an SVG to 250x250 in the browser, clipped to a circle like TradingView shows it."""
    from PIL import Image

    svg = re.sub(r"<svg\b", f'<svg width="{SIZE}" height="{SIZE}"', svg, count=1)
    radius = "0" if square else "50%"
    page.set_content(
        "<html><body style='margin:0;background:transparent'>"
        f"<div style='width:{SIZE}px;height:{SIZE}px;border-radius:{radius};overflow:hidden'>"
        f"{svg}</div></body></html>")
    png = page.screenshot(omit_background=True, clip={"x": 0, "y": 0, "width": SIZE, "height": SIZE})
    return Image.open(io.BytesIO(png))


def cmd_tradingview(args):
    from playwright.sync_api import sync_playwright

    rows = load_sources()
    log = load_log(args.own)
    todo = [r for r in rows
            if args.overwrite or not log.get((r["venue"], r["symbol"]), "").startswith("tradingview")]
    if args.keep_existing:
        todo = [r for r in todo if not os.path.exists(own_path(args.own, r["venue"], r["symbol"]))]
    print(f"Looking up {len(todo)} TradingView logo ids ...")
    logoids = {}
    for i in range(0, len(todo), 100):
        logoids.update(tradingview_logoids([r["tradingview"] for r in todo[i:i + 100]]))

    missing = []
    with sync_playwright() as p:
        try:
            browser = p.chromium.launch(channel="msedge")
        except Exception:
            browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": SIZE, "height": SIZE}, device_scale_factor=1)
        for r in todo:
            key = (r["venue"], r["symbol"])
            logoid = logoids.get(r["tradingview"]) or tradingview_search_logoid(r["tradingview"])
            if not logoid or logoid.startswith("country/"):
                missing.append(r)
                continue
            req = urllib.request.Request(LOGO_URL.format(logoid), headers=UA)
            try:
                with urllib.request.urlopen(req, timeout=30) as resp:
                    svg = resp.read().decode("utf-8")
            except Exception as e:
                print(f"  {r['tradingview']}: download failed ({e})")
                missing.append(r)
                continue
            save_png(render_svg(page, svg, args.square), own_path(args.own, *key))
            log[key] = f"tradingview:{logoid}"
            print(f"  {r['tradingview']:<22} -> {logoid}")
        browser.close()
    save_log(args.own, log)
    print(f"Done. {len(todo) - len(missing)} downloaded, {len(missing)} without a TradingView logo:")
    for r in missing:
        print(f"  {r['tradingview']}  {r['name']}")
    cmd_placeholders(args)
    cmd_report(args)


# ---------------------------------------------------------------- placeholders (last resort)

FONTS = ["arialbd.ttf", "C:/Windows/Fonts/arialbd.ttf", "DejaVuSans-Bold.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "/Library/Fonts/Arial Bold.ttf"]
PALETTE = ["#2962FF", "#00897B", "#7B1FA2", "#C62828", "#EF6C00", "#2E7D32", "#AD1457", "#0277BD",
           "#5D4037", "#455A64", "#6A1B9A", "#00838F"]
STOP = {"AND", "OF", "THE", "CO", "GROUP", "HOLDINGS", "INC", "LTD", "PLC", "SA", "AG", "&"}


def initials(name):
    words = [w for w in re.split(r"[\s\-/()]+", name) if w and w.upper() not in STOP]
    if not words:
        return "?"
    if len(words) == 1:
        w = words[0]
        return w.upper() if w.isupper() and len(w) <= 4 else w[0].upper()
    return (words[0][0] + words[1][0]).upper()


def placeholder(name):
    import zlib
    from PIL import Image, ImageDraw, ImageFont

    scale = 4
    s = SIZE * scale
    img = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.ellipse((0, 0, s - 1, s - 1), fill=PALETTE[zlib.crc32(name.encode()) % len(PALETTE)])
    text = initials(name)
    font = None
    for f in FONTS:
        try:
            font = ImageFont.truetype(f, int(s * {1: 0.5, 2: 0.42, 3: 0.32}.get(len(text), 0.25)))
            break
        except OSError:
            pass
    font = font or ImageFont.load_default(size=int(s * 0.4))
    d.text((s / 2, s / 2), text, font=font, fill="white", anchor="mm")
    return img.resize((SIZE, SIZE), Image.LANCZOS)


def cmd_placeholders(args):
    log = load_log(args.own)
    made = 0
    for row in load_sources():
        key = (row["venue"], row["symbol"])
        path = own_path(args.own, *key)
        if os.path.exists(path):
            continue
        save_png(placeholder(row["name"]), path)
        log[key] = "placeholder"
        made += 1
    save_log(args.own, log)
    print(f"{made} placeholders written")


# ---------------------------------------------------------------- report

def cmd_report(args):
    import openpyxl

    nv = set(open(NVSTLY_INDEX, encoding="utf-8").read().split())
    names = {(r["venue"], r["symbol"]): r["name"] for r in load_sources()}
    log = load_log(args.own)
    wb = openpyxl.load_workbook(os.path.join(ROOT, "instrument_catalogue.xlsx"), read_only=True)
    rows = list(wb["Catalogue"].iter_rows(values_only=True))[1:]
    out = []
    for r in rows:
        symbol, tv, venue, group = r[1], r[3], r[6], r[12] or ""
        key = (venue, symbol)
        if os.path.exists(own_path(args.own, venue, symbol)):
            source = log.get(key, "own")
            status = "placeholder" if source == "placeholder" else "added"
        elif venue in US_VENUES and (symbol in nv or symbol.replace(".", "-") in nv):
            status, source = "have", "nvstly"
        else:
            status, source = "missing", ""
        out.append({"status": status, "venue": venue, "symbol": symbol, "tradingview": tv,
                    "name": names.get(key, ""), "source": source,
                    "file": f"own/{venue}/{symbol.replace(':', '_')}.png" if source and status != "have" else "",
                    "group": group})
    order = {"missing": 0, "placeholder": 1, "added": 2, "have": 3}
    out.sort(key=lambda o: (order[o["status"]], o["venue"], o["symbol"]))

    os.makedirs(os.path.join(ROOT, "reports"), exist_ok=True)
    with open(os.path.join(ROOT, "reports", "logo_status.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(out[0]))
        w.writeheader()
        w.writerows(out)

    count = {s: sum(o["status"] == s for o in out) for s in order}
    titles = {"missing": "Still missing", "placeholder": "Lettered placeholder in own/ (no real logo found yet)",
              "added": "Real logo added in own/", "have": "Already covered by nvstly"}
    md = ["# Logo status", "",
          f"{len(out)} instruments in instrument_catalogue.xlsx: "
          f"{count['have']} already covered by nvstly, {count['added']} real logos added in own/, "
          f"{count['placeholder']} lettered placeholders, {count['missing']} still missing.", ""]
    for s in order:
        md += [f"## {titles[s]} ({count[s]})", "", "| Venue | Ticker | TradingView | Company | Source |",
               "|---|---|---|---|---|"]
        md += [f"| {o['venue']} | {o['symbol']} | {o['tradingview']} | {o['name']} | {o['source']} |"
               for o in out if o["status"] == s]
        md.append("")
    with open(os.path.join(ROOT, "reports", "logo_status.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md))
    print(", ".join(f"{k} {v}" for k, v in count.items()))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--own", default=OWN, help="output folder (default: own/ in this repo)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("--nvstly", required=True, help="checkout of github.com/nvstly/icons")
    b.add_argument("--simple-icons", required=True, help="checkout of github.com/simple-icons/simple-icons")
    b.add_argument("--overwrite", action="store_true")
    t = sub.add_parser("tradingview")
    t.add_argument("--overwrite", action="store_true", help="also redo logos already fetched from TradingView")
    t.add_argument("--keep-existing", action="store_true", help="only fill tickers that have no file yet")
    t.add_argument("--square", action="store_true", help="keep the square logo instead of TradingView's circle")
    sub.add_parser("placeholders")
    sub.add_parser("report")
    args = ap.parse_args()
    {"build": cmd_build, "tradingview": cmd_tradingview, "placeholders": cmd_placeholders, "report": cmd_report}[args.cmd](args)


if __name__ == "__main__":
    sys.exit(main())
